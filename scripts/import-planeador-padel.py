"""
Importa el horario de la ACADEMIA DE PÁDEL (24-sep-2026).

    npm run import:padel -- "ruta/HORARIO ACADEMIA DE PADEL LEO.xlsx"            # simulacro
    npm run import:padel -- "ruta/HORARIO ACADEMIA DE PADEL LEO.xlsx" --apply    # escribe

Misma lógica que tenis (la clase es profesor + día + hora + duración; la
academia —recreativa o competencia— es de cada niño), pero el Excel es otro:
una sola rejilla, sin hoja de base de datos ni documentos. Cada fila es
"HORA + PROFESOR" y cada celda trae los nombres separados por saltos de línea o
por varios espacios.

Todo lo que el Excel no dice, o dice mal, va en las listas EXPLÍCITAS de abajo,
cada una con quién lo decidió. Nada se adivina: esto decide a quién se le cobra.

⚠️ Solo toca PÁDEL: clases de pádel y matrículas de las academias de pádel. El
importador de tenis, a su vez, solo toca tenis.
"""
import openpyxl, json, urllib.request, urllib.error, sys, os, unicodedata, re, warnings, datetime

warnings.filterwarnings("ignore")
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_args = [a for a in sys.argv[1:] if not a.startswith("--")]
if not _args:
    sys.exit("Falta la ruta del Excel.")
XLSX = _args[0]
APPLY = "--apply" in sys.argv
VIGENTE_DESDE = "2026-10-01"  # igual que tenis: el planeador aplica desde el 1-oct (Laura)

env = {}
for line in open(REPO + "/.env"):
    line = line.strip()
    if "=" in line and not line.startswith("#"):
        k, v = line.split("=", 1); env[k] = v.strip().strip('"')
URL = env["NEXT_PUBLIC_SUPABASE_URL"]; KEY = env["SUPABASE_SERVICE_ROLE_KEY"]
HDR = {"apikey": KEY, "Authorization": f"Bearer {KEY}", "Content-Type": "application/json"}

def req(method, path, body=None, prefer=None):
    h = dict(HDR)
    if prefer: h["Prefer"] = prefer
    data = json.dumps(body).encode() if body is not None else None
    r = urllib.request.Request(URL + "/rest/v1/" + path, data=data, headers=h, method=method)
    try:
        with urllib.request.urlopen(r) as resp:
            raw = resp.read(); return json.loads(raw) if raw else None
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"{method} {path[:70]} -> {e.code} {e.read().decode()[:300]}")

def get(path): return req("GET", path)
def insert(table, rows): return req("POST", table, rows, "return=representation")
def patch(path, body): return req("PATCH", path, body)

def nm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()

# ── Decisiones explícitas ───────────────────────────────────────────────────
PROFESOR = {"LEO": "Leo Ruíz", "JUAN": "Juan Cruz", "VICTOR": "Victor Acosta"}

# Clases de colegio: van como clase, SIN niños (Laura, 24-sep-2026).
COLEGIOS = {"MONTESSORI": "Montessori"}

# Dos nombres pegados en una misma línea del Excel.
PARTIR = {"SALVADOR OLARTE JHONTAN DULCEY": ["SALVADOR OLARTE", "JHONTAN DULCEY"]}

# Lo que NO es academia.
PARTICULAR = re.compile(r"\(\s*PARTICULAR\s*\)")  # "VALENTINO GOMEZ (Particular)": clase particular

# Quitar un niño de una celda donde el Excel lo repite por error.
# Isaak Salgado salía el jueves 4 p. m. con Juan Y con Leo: va con Leo (Laura, 24-sep).
QUITAR = {("ISAAK SALGADO", "JUAN", 4, "16:00")}

# Los ÚNICOS de competencia: lunes y miércoles 5–6:30 p. m. con Leo (Laura, 24-sep).
COMPETENCIA = {"AGUSTIN PEREZ", "MATIAS GARCIA", "EMILIO OLARTE", "PEDRO CORREA"}

# Nombre del Excel → ficha (miembro) real, cuando no casa tal cual. Con su evidencia.
MIEMBRO = {
    "NICOLE BUSTAMANTE": 468,     # "Nicol Bustamante" (Laura)
    "JULIA VELEZ JIMENEZ": 481,   # Julia Vélez; mamá María Paulina Jiménez (Laura)
    "PEDRO GOMEZ LASERNA": 408,   # Pedro Gómez; mamá Aleja Serna Gómez (Laura)
    "SOFIA MORENO": 477,          # Sophie Moreno, hija de Pierre Moreno (Laura)
    "SALVADOR OLARTE": 474,       # hay dos fichas; esta tiene documento (la 324 queda por revisar)
}

# Niños que no existían: se crean como hermanos dentro de la ficha de la familia.
CREAR = {
    "ANA BARBERA": {"cliente_id": 58, "nombres": "Ana", "apellidos": "Barbera"},      # hermana de Armando Barbera (Laura)
    "JHONTAN DULCEY": {"cliente_id": 37, "nombres": "Jhontan", "apellidos": "Dulcey"},  # hijo de Jhon Dulcey (Laura)
}

# Sin resolver todavía: NO se cargan y se reportan.
PENDIENTES = {"SIMON MEJIA", "SALOMON AGUDELO"}

DIAS = {"LUNES": 1, "MARTES": 2, "MIERCOLES": 3, "JUEVES": 4, "VIERNES": 5, "SABADO": 6}

def hora24(txt):
    """'4:00 p. m.' → '16:00'. Todas las del Excel son de la tarde."""
    m = re.search(r"(\d{1,2}):(\d{2})\s*([AP])?", txt)
    h, mi, ap = int(m.group(1)), m.group(2), m.group(3)
    if ap == "P" and h < 12: h += 12
    return f"{h:02d}:{mi}", m.end()

# ── 1 · Leer la rejilla ─────────────────────────────────────────────────────
ws = openpyxl.load_workbook(XLSX, data_only=True).active
enc = [nm(c.value) for c in ws[2]]
col_dia = {i: DIAS[v] for i, v in enumerate(enc) if v in DIAS}

clases = {}      # (profesor, dia, hora) → {"dur": min, "colegio": str|None, "ninos": [..]}
avisos = []
for row in ws.iter_rows(min_row=3, values_only=True):
    cab = nm(row[0]).replace(".", "").replace("\xa0", " ")
    if not cab: continue
    ini, fin = hora24(cab)
    resto = cab[fin:]
    dur = 60  # Laura: todo es de 60 min salvo el grupo de 5 a 6:30
    m2 = re.search(r"\bA\s+(\d{1,2}:\d{2}\s*[AP]?)", resto)
    if m2:
        f, _ = hora24(m2.group(1))
        dur = (int(f[:2]) * 60 + int(f[3:])) - (int(ini[:2]) * 60 + int(ini[3:]))
    prof = next((k for k in PROFESOR if re.search(rf"\b{k}\b", cab)), None)
    if not prof:
        avisos.append(f"Fila sin profesor reconocible: {cab!r}"); continue
    for i, dia in col_dia.items():
        v = row[i] if i < len(row) else None
        if not v or not str(v).strip(): continue
        nombres = []
        for x in re.split(r"\n|\s{3,}", str(v)):
            x = nm(x)
            if x: nombres += PARTIR.get(x, [x])
        clave = (prof, dia, ini)
        c = clases.setdefault(clave, {"dur": dur, "colegio": None, "ninos": []})
        for x in nombres:
            if x in COLEGIOS:
                c["colegio"] = COLEGIOS[x]
            elif PARTICULAR.search(x):
                avisos.append(f"Se omite (clase particular, no academia): {x} · {prof} día {dia} {ini}")
            elif (x, prof, dia, ini) in QUITAR:
                avisos.append(f"Se quita (decisión del club): {x} · {prof} día {dia} {ini}")
            else:
                c["ninos"].append(x)

# Filas del Excel sin ninguna celda (Victor a las 3 y a las 4) no crean clase, y
# tampoco una celda que solo traía una particular (lunes 5 p. m. de Juan).
for k in [k for k, c in clases.items() if not c["ninos"] and not c["colegio"]]:
    del clases[k]
print(f"Clases en el Excel: {len(clases)}")
for (p, d, h), c in sorted(clases.items(), key=lambda kv: (kv[0][1], kv[0][2], kv[0][0])):
    print(f"  {['','Lun','Mar','Mié','Jue','Vie','Sáb'][d]} {h} {PROFESOR[p]:14} {c['dur']:>3} min · "
          + (f"COLEGIO {c['colegio']}" if c["colegio"] else f"{len(c['ninos'])} niños"))

# ── 2 · Cruzar niños con las fichas ─────────────────────────────────────────
miembros = get("cliente_miembros?select=id,cliente_id,nombres,apellidos,activo&limit=5000")
por_id = {m["id"]: m for m in miembros}
def buscar(nombre):
    if nombre in MIEMBRO: return por_id.get(MIEMBRO[nombre])
    t = nombre.split()
    hits = [m for m in miembros if m["activo"] and all(x in nm(f"{m['nombres']} {m['apellidos']}").split() for x in t)]
    return hits[0] if len(hits) == 1 else ("AMBIGUO", hits) if hits else None

todos = sorted({n for c in clases.values() for n in c["ninos"]})
ninos, crear, pendientes = {}, [], []
for n in todos:
    if n in PENDIENTES: pendientes.append(n); continue
    if n in CREAR:
        c = CREAR[n]
        ya = [m for m in miembros if m["cliente_id"] == c["cliente_id"] and nm(m["nombres"]) == nm(c["nombres"]) and nm(m["apellidos"]) == nm(c["apellidos"])]
        if ya: ninos[n] = ya[0]
        else: crear.append(n)
        continue
    m = buscar(n)
    if isinstance(m, tuple):
        avisos.append(f"AMBIGUO {n}: " + ", ".join(f"m{x['id']} {x['nombres']} {x['apellidos']}" for x in m[1]))
    elif m: ninos[n] = m
    else: pendientes.append(n)

print(f"\nNiños: {len(todos)} · cruzan {len(ninos)} · se crean {len(crear)} · pendientes {len(pendientes)}")
for n in crear: print(f"  + crear {n} en la ficha {CREAR[n]['cliente_id']}")
for n in pendientes: print(f"  ⚠️ sin cargar (no existe): {n}")
for a in avisos: print(f"  · {a}")
if any(a.startswith("AMBIGUO") for a in avisos):
    sys.exit("\nHay nombres ambiguos: agrégalos a MIEMBRO antes de seguir.")

profes = {p["nombre"]: p["id"] for p in get("profiles?select=id,nombre")}
prof_id = {k: profes[v] for k, v in PROFESOR.items()}
acas = get("academias?select=id,deporte,categoria")
ACA = {a["categoria"]: a["id"] for a in acas if a["deporte"] == "padel"}

cupos = sum(len(c["ninos"]) for c in clases.values())
print(f"\nA escribir: {len(clases)} clases · {len(ninos) + len(crear)} matrículas · {cupos} cupos"
      f" (vigentes desde {VIGENTE_DESDE})")
if not APPLY:
    print("\n(simulacro — nada se escribió. Agrega --apply para aplicar)")
    sys.exit(0)

# ── 3 · Escribir (idempotente, solo pádel) ──────────────────────────────────
for n in crear:
    c = CREAR[n]
    ninos[n] = insert("cliente_miembros", [dict(cliente_id=c["cliente_id"], nombres=c["nombres"],
                                                 apellidos=c["apellidos"], deportes=["padel"],
                                                 es_titular=False, activo=True)])[0]
print(f"cliente_miembros: +{len(crear)}")

ya = {(c["profesor_id"], c["dia_semana"], c["hora_inicio"][:5]): c
      for c in get("clase_semanal?select=id,profesor_id,dia_semana,hora_inicio,duracion_min,colegio&activa=eq.true&deporte=eq.padel")}
nuevas, cambiadas = [], 0
for (p, d, h), c in clases.items():
    k = (prof_id[p], d, h)
    if k in ya:
        x = ya[k]
        if x["duracion_min"] != c["dur"] or x["colegio"] != c["colegio"]:
            patch(f"clase_semanal?id=eq.{x['id']}", {"duracion_min": c["dur"], "colegio": c["colegio"]}); cambiadas += 1
    else:
        nuevas.append(dict(profesor_id=prof_id[p], deporte="padel", dia_semana=d, hora_inicio=f"{h}:00",
                           duracion_min=c["dur"], colegio=c["colegio"], vigente_desde=VIGENTE_DESDE))
for c in insert("clase_semanal", nuevas) if nuevas else []:
    ya[(c["profesor_id"], c["dia_semana"], c["hora_inicio"][:5])] = c
print(f"clase_semanal: +{len(nuevas)} · {cambiadas} ajustadas (total {len(ya)})")

quiere = {ninos[n]["id"]: ACA["competencia"] if n in COMPETENCIA else ACA["recreativa"] for n in ninos}
vivas = get("inscripciones?select=id,miembro_id,academia_id,activa&academia_id=in.(" + ",".join(str(v) for v in ACA.values()) + ")")
insc = {}
for i in vivas:
    if i["miembro_id"] in quiere and (i["activa"] or i["miembro_id"] not in insc): insc[i["miembro_id"]] = i
familia = {m["id"]: m["cliente_id"] for m in ninos.values()}
faltan = [dict(academia_id=a, cliente_id=familia[mid], miembro_id=mid)
          for mid, a in quiere.items() if mid not in insc]
for i in insert("inscripciones", faltan) if faltan else []:
    insc[i["miembro_id"]] = i
movidas = 0
for mid, a in quiere.items():
    i = insc[mid]
    if i["academia_id"] != a or not i.get("activa", True):
        patch(f"inscripciones?id=eq.{i['id']}", {"academia_id": a, "activa": True, "retirada_el": None}); movidas += 1
# Al que ya no está en el horario se le RETIRA (no se borra).
fuera = [i for i in vivas if i["activa"] and i["miembro_id"] not in quiere]
for i in fuera:
    patch(f"inscripciones?id=eq.{i['id']}", {"activa": False, "retirada_el": datetime.date.today().isoformat()})
print(f"inscripciones: +{len(faltan)} · {movidas} ajustadas · {len(fuera)} retiradas")

clases_padel = {c["id"] for c in ya.values()}
links_ya = {(l["inscripcion_id"], l["clase_id"]) for l in get("inscripcion_clase?select=inscripcion_id,clase_id&limit=5000")
            if l["clase_id"] in clases_padel}
quiero = set()
for (p, d, h), c in clases.items():
    cid = ya[(prof_id[p], d, h)]["id"]
    for n in c["ninos"]:
        if n in ninos: quiero.add((insc[ninos[n]["id"]]["id"], cid))
add = [dict(inscripcion_id=i, clase_id=c) for i, c in quiero - links_ya]
if add: insert("inscripcion_clase", add)
sobran = links_ya - quiero
for i, c in sobran:
    req("DELETE", f"inscripcion_clase?inscripcion_id=eq.{i}&clase_id=eq.{c}")
print(f"inscripcion_clase: +{len(add)} · -{len(sobran)} (total {len(quiero)})")
print("\n✅ Listo.")
